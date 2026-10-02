import json, os, re, subprocess, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

S = os.path.dirname(os.path.abspath(__file__))
subprocess.run([sys.executable, os.path.join(S, "merge.py")], check=True, stdout=subprocess.DEVNULL)
M = json.load(open(os.path.join(S, "merged.json")))
rows, checks, udia = M["rows"], M["checks"], M["udia_all"]
OUT = sys.argv[1]

DISPLAY = {"AHB Group": "AHB Group (now Mirrastone)", "Costa Asset Management": "Costa Property Group / Costa Asset Management",
           "Exford Waters": "Exford Waters (Wegg Pty Ltd)", "URBAN Land Projects": "Urban Land Projects",
           "XWISE Group": "XWISE Group"}
MSG_NAME = {"AHB Group": "Mirrastone", "Costa Asset Management": "Costa Property Group", "Exford Waters": "the Exford Waters team",
            "The Corcoris Group Developments Pty Ltd": "The Corcoris Group", "URBAN Land Projects": "Urban Land Projects",
            "Star Investment Group Australia (Star Marketing)": "Star Investment Group", "Gull & Company": "Gull & Co",
            "Wel.Co": "Wel.Co", "GURNER": "GURNER", "ID_Land": "ID_Land"}
SKIP = {  # left out of the Outreach tab, with the reason
    "Development Victoria": "Government agency", "Homes Victoria": "Government agency",
    "Community Housing Victoria Ltd": "Community housing provider", "Housing Choices Australia Limited (Victoria)": "Community housing provider",
    "HousingFirst Ltd": "Community housing provider", "Launch Housing": "Homelessness charity", "Bolton Clarke": "Aged-care not-for-profit",
    "Banner Asset Management": "Finance firm", "Three Thirds Group": "Professional services firm",
    "Antipodean Land Developments": "Only contact is a personal gmail", "Sunrise Ventures": "No public contact details",
    "The Range (Trafalgar) Pty Ltd": "Landowner entity, no contacts", "Trafalgar Property Developments Pty Ltd": "Landowner entity, no contacts",
    "Narracan Meadows developer (entity not identified)": "Developer not identified; only agent contact",
    "Hunter Park Country Estate developer (entity not identified)": "Developer not identified; only agent contact",
    "Cambrooke Property Pty Ltd": "Only selling-agent contact",
    "LOGOS": "Site now redirects to ESR Australia — contact ESR instead",
    "Mount Atkinson Holdings": "Domain no longer belongs to the company",
    "Eight Property Investments": "In liquidation (ASIC notice)", "Montego Homes": "Collapsed in 2024",
    "AVJennings": "Website now redirects to AVID Property Group — contact AVID instead", "Seebeck Group Enterprise": "No contact details",
}
MARKETS = ["Geelong", "Bendigo", "Ballarat", "Shepparton", "Kilmore", "Beveridge", "Moe", "Wollert", "Tarneit", "Truganina",
           "Werribee", "Melton", "Wyndham", "Mickleham", "Craigieburn", "Kalkallo", "Donnybrook", "Lara", "Armstrong Creek",
           "Drysdale", "Clyde", "Officer", "Pakenham", "Cranbourne", "Newborough", "Kilmore", "Wallan", "Sunbury"]

def msg_name(dev):
    if dev in MSG_NAME: return MSG_NAME[dev]
    return re.sub(r"\s+(Pty\.? Ltd\.?|Limited|Ltd)$", "", dev).strip()

def area(projects):
    found = []
    for m in MARKETS:
        if re.search(r"\b" + m + r"\b", projects) and m not in found: found.append(m)
    if not found: return ""
    return found[0] if len(found) == 1 else ", ".join(found[:-1][:3]) + " and " + found[min(len(found) - 1, 3)]

BODY = """Hi Acquisition Team,

My name is Inder Sandhu with Crown Real Estate Agents. I specialise in development land across Melbourne's western and northern growth corridors and regional Victoria.

I have development sites available across Wollert, Beveridge, Tarneit, Geelong, Deanside, Bonnie Brook, Fraser Rise and Truganina:

* Shovel-ready residential and industrial.
* Permit-approved townhouse sites, including one for 44 townhouses.
* Childcare and townhouse sites, both raw and approved.
* About 50 acres of investigation-area land.

Given {who} active presence {where}, I believe these sites would be a strong fit for your acquisitions pipeline.
Could you please direct me to the appropriate contact in your acquisitions department so I can share further details?

If you'd rather not hear from me about sites, just let me know and I won't contact you again.

Kind regards,
Inder Sandhu
Crown Real Estate Agents
0484 926 324 | inder@crownrea.com.au"""

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10)
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")

def sheet(wb, title, headers, data, widths, wrap_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in data: ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.font = BODY_FONT
            c.alignment = Alignment(wrap_text=c.column in wrap_cols, vertical="top")
            if isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink = c.value
                c.font = Font(name=FONT, size=10, color="0563C1", underline="single")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32
    return ws

wb = Workbook(); wb.remove(wb.active)

# 1. Main list
cols = list(rows[0].keys())
data = []
for i, r in enumerate(rows, 1):
    v = [r[c] for c in cols]
    v[0] = DISPLAY.get(r["Developer"], r["Developer"])
    data.append([i] + v)
widths = [5, 30, 26, 32, 15, 34, 26, 30, 30, 28, 48, 11, 28, 26, 8, 8, 8, 8, 8, 26, 50, 60]
ws = sheet(wb, "Developer List", ["#"] + cols, data, widths, wrap_cols=(3, 11, 13, 21, 22))
fc = cols.index("Cross-check flags") + 2
for row in ws.iter_rows(min_row=2):
    if row[fc - 1].value: row[fc - 1].fill = FLAG_FILL

# 2. Outreach
out, skipped = [], []
for r in rows:
    dev = r["Developer"]
    if dev in SKIP: skipped.append([DISPLAY.get(dev, dev), SKIP[dev]]); continue
    a = area(r["Regional projects / footprint"])
    where = f"in {a}" if a else "across Melbourne's growth corridors and regional Victoria"
    form = r["Contact page"] if "openlot.com.au" not in r["Contact page"] else ""
    form = form or r["Website"] or "No website — use email / phone"
    out.append([DISPLAY.get(dev, dev), form, r["Website"], r["Email"] if "@" in r["Email"] else "", r["Phone"],
                "Inder Sandhu", "inder@crownrea.com.au", "0484926324", "Crown Real Estate Agents",
                "Development sites – Wollert, Beveridge, Tarneit, Geelong & more",
                BODY.format(who=(lambda n: n + ("'" if n.endswith('s') else "'s"))(msg_name(dev)), where=where), "Not sent", "", ""])
oh = ["Developer", "Contact form / contact page", "Website", "Published email", "Phone", "Your name", "Your email",
      "Your mobile", "Company", "Subject", "Message (personalised)", "Status", "Date sent", "Response / notes"]
ws = sheet(wb, "Outreach", oh, out, [30, 40, 30, 32, 15, 13, 24, 12, 22, 34, 80, 11, 11, 30], wrap_cols=(11,))
for row in ws.iter_rows(min_row=2):
    ws.row_dimensions[row[0].row].height = 150
ws.cell(row=len(out) + 3, column=1, value="Left out of outreach (and why):").font = Font(name=FONT, bold=True)
for j, (n, why) in enumerate(skipped):
    ws.cell(row=len(out) + 4 + j, column=1, value=n).font = BODY_FONT
    ws.cell(row=len(out) + 4 + j, column=2, value=why).font = BODY_FONT

# 3. Cross-check
sheet(wb, "Cross-check", ["Developer", "Differences between sources (your workbook vs UDIA vs fresh research)"],
      [[DISPLAY.get(c["Developer"], c["Developer"]), c["Flags"]] for c in checks], [34, 120], wrap_cols=(2,))

# 4. UDIA – every member
sheet(wb, "UDIA All Members", ["Page", "Name", "Category", "Location", "Phone", "Email", "Specialties", "Website"],
      [[m["page"], m["name"], m["category"], m["location"], m["phone"], m["email"], m["specialties"], m["website"]] for m in udia],
      [6, 36, 24, 22, 15, 32, 50, 36], wrap_cols=(7,))

# 5. Sources
notes = [
    ["Your developer list (49 names)", "Pasted in chat."],
    ["Regional table (Geelong, Bendigo, Ballarat, Shepparton, Kilmore, Beveridge, Moe)", "Pasted in chat. 'Private landowners/developers – Moe' researched separately (column 'Moe private developer')."],
    ["Your workbook 16872595-Developer_List_.xlsx", "Developer Universe, Regional Projects and Decision Makers tabs cross-checked; every developer included."],
    ["UDIA Victoria members directory", "All 34 pages (https://udiavic.com.au/members-directory/) scraped 2 Oct 2026. Every member is on the 'UDIA All Members' tab; members with category Developer or a Developer specialty are in the main list."],
    ["Fresh web research, 2 Oct 2026", "Company contact pages, estate sites, directories. Field sources column shows which source each website/phone/email came from (Research > UDIA > Workbook)."],
    ["Email rule", "Only emails actually published somewhere are listed. None were guessed. 'Contact form only' means the company publishes no email."],
    ["Two different 'North Quarter' estates", "The North Quarter, North Shepparton = Urban Land Projects (thenorthquarter.com.au). North Quarter, Newborough/Moe = Solovey (north-quarter.com.au)."],
    ["Lendlease", "Sold its Victorian communities (Atherstone, Aurora, Averley, Harpley) to Stockland and Supalai in Nov 2024."],
    ["Star Investment Group (Lake Narracan)", "Investor concerns about this scheme are reported online — check before engaging."],
    ["Outreach tab", "Message is personalised per developer (name + area). Added an opt-out line, as the Spam Act requires a way to opt out of commercial messages. Fill Status / Date sent as you go."],
]
sheet(wb, "Sources & Notes", ["Item", "Detail"], notes, [55, 120], wrap_cols=(1, 2))

wb.save(OUT)
print("saved", OUT, len(rows), "developers;", len(out), "outreach rows;", len(skipped), "skipped;", len(checks), "flags")
